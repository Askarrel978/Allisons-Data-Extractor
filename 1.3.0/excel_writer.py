"""Excel output helpers for Allison's Data Extractor.

Adds automatic backups and duplicate detection while preserving the original
six-column workbook format. Values are written into the workbook's existing
columns by header name, so other column orders and extra columns are left
alone, and the workbook is saved through a temporary file so an interrupted
save cannot corrupt the original.
"""

import os
import re
import shutil
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from pdf_extractor import HEADERS


def _normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _header_columns(ws) -> dict:
    """Map each of HEADERS to its 1-based column number in row 1, when present."""
    wanted = {_normalize_header(header): header for header in HEADERS}
    found = {}
    for column, value in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()), start=1):
        header = wanted.get(_normalize_header(value))
        if header and header not in found:
            found[header] = column
    return found


def _select_sheet(wb):
    """Return the sheet holding the extractor's data rather than whichever sheet is active."""
    for ws in wb.worksheets:
        if len(_header_columns(ws)) * 2 >= len(HEADERS):
            return ws
    return wb.active


def _last_header_column(ws) -> int:
    last = 0
    for column, value in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()), start=1):
        if value not in (None, ""):
            last = column
    return last


def _data_columns(ws) -> dict:
    """Columns to read/write each field in.

    A header row that matches none of our names probably just labels them
    differently, so it keeps the original fixed A-F layout instead.
    """
    columns = _header_columns(ws)
    if not columns and _last_header_column(ws):
        return {header: index for index, header in enumerate(HEADERS, start=1)}
    return columns


def _last_data_row(ws, columns) -> int:
    """Last row with a value in any of ``columns``; styled but empty rows are ignored."""
    columns = sorted(columns)
    if not columns:
        return 1
    last = 1
    for row_number, values in enumerate(
        ws.iter_rows(min_row=1, min_col=columns[0], max_col=columns[-1], values_only=True), start=1
    ):
        if any(values[col - columns[0]] not in (None, "") for col in columns if col - columns[0] < len(values)):
            last = row_number
    return last


def _autosize_columns(ws, columns):
    for column in columns:
        max_len = 0
        for (value,) in ws.iter_rows(min_col=column, max_col=column, values_only=True):
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(column)].width = min(max_len + 2, 60)


def create_backup(filepath: str) -> str:
    """Create a timestamped backup beside an existing workbook."""
    if not os.path.exists(filepath):
        return ""
    folder = os.path.dirname(os.path.abspath(filepath))
    backup_dir = os.path.join(folder, "Allisons Data Extractor Backups")
    os.makedirs(backup_dir, exist_ok=True)
    stem, ext = os.path.splitext(os.path.basename(filepath))
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"{stem}_backup_{stamp}{ext or '.xlsx'}")
    shutil.copy2(filepath, backup_path)
    return backup_path


def _ensure_not_locked(filepath: str) -> None:
    """Raise PermissionError if another program (usually Excel) has the file open."""
    with open(filepath, "r+b"):
        pass


def _temp_save_path(filepath: str) -> str:
    folder = os.path.dirname(os.path.abspath(filepath))
    stem, ext = os.path.splitext(os.path.basename(filepath))
    return os.path.join(folder, f".{stem}.saving{ext or '.xlsx'}")


def _save_atomically(wb, filepath: str) -> None:
    """Save to a temporary file and then swap it into place in one step."""
    temp_path = _temp_save_path(filepath)
    try:
        wb.save(temp_path)
        os.replace(temp_path, filepath)
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


def _existing_rows(filepath: str):
    if not os.path.exists(filepath):
        return []
    wb = load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws = _select_sheet(wb)
        columns = _data_columns(ws)
        rows = []
        for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            item = {
                header: values[columns[header] - 1] if header in columns and columns[header] - 1 < len(values) else ""
                for header in HEADERS
            }
            item["_excel_row"] = excel_row
            rows.append(item)
        return rows
    finally:
        wb.close()


def known_values(filepath: str, field: str) -> set:
    """All non-empty values of one field already in the workbook (e.g. known lot codes)."""
    try:
        rows = _existing_rows(filepath)
    except Exception:
        # A missing, locked or unreadable workbook simply means nothing is known.
        return set()
    return {str(row.get(field) or "").strip() for row in rows} - {""}


def find_duplicates(filepath: str, rows: list) -> list:
    """Find likely duplicate shipments based primarily on Sales Order Number."""
    existing = _existing_rows(filepath)
    duplicates = []
    seen_incoming = {}
    for incoming_index, incoming in enumerate(rows):
        so = str(incoming.get("Sales Order Number", "") or "").strip()
        lot = str(incoming.get("Lot Code", "") or "").strip().upper()
        if not so:
            continue
        if so in seen_incoming:
            duplicates.append({
                "incoming_index": incoming_index,
                "sales_order": so,
                "lot_code": lot,
                "excel_row": f"incoming row {seen_incoming[so] + 1}",
                "match_type": "Duplicate within current batch",
            })
        else:
            seen_incoming[so] = incoming_index
        for old in existing:
            old_so = str(old.get("Sales Order Number", "") or "").strip()
            old_lot = str(old.get("Lot Code", "") or "").strip().upper()
            if old_so != so:
                continue
            match_type = "Sales Order Number"
            if lot and old_lot and lot == old_lot:
                match_type = "Sales Order Number + Lot Code"
            duplicates.append({
                "incoming_index": incoming_index,
                "sales_order": so,
                "lot_code": lot,
                "excel_row": old.get("_excel_row"),
                "match_type": match_type,
            })
            break
    return duplicates


def append_rows(filepath: str, rows: list, create_backup_first: bool = True):
    """Append rows and return the backup path created before writing, if any."""
    exists = os.path.exists(filepath)
    if exists:
        _ensure_not_locked(filepath)
    backup_path = create_backup(filepath) if create_backup_first else ""

    if exists:
        wb = load_workbook(filepath)
        ws = _select_sheet(wb)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"

    # Match the workbook's own header row. Any of our headers it lacks are
    # added to the right of its existing columns (all six for a blank sheet).
    columns = _data_columns(ws)
    added_columns = []
    next_column = _last_header_column(ws) + 1
    for header in HEADERS:
        if header in columns:
            continue
        cell = ws.cell(row=1, column=next_column, value=header)
        cell.font = Font(bold=True)
        columns[header] = next_column
        added_columns.append(next_column)
        next_column += 1

    start_row = _last_data_row(ws, columns.values()) + 1
    for offset, row in enumerate(rows):
        for header, column in columns.items():
            value = row.get(header, "")
            ws.cell(row=start_row + offset, column=column, value=value if value != "" else None)

    # Only size columns this export created; the user's own widths are kept.
    _autosize_columns(ws, added_columns)
    _save_atomically(wb, filepath)
    return backup_path
