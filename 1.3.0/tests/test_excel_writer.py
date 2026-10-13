import os
import stat

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side

from excel_writer import append_rows, find_duplicates, known_values
from pdf_extractor import HEADERS


def _row(so="123456", lot="L-77", **overrides):
    row = {
        "Date": "August 3, 2026",
        "Sales Order Number": so,
        "Customer and Location": "Acme Farms, 1 Main St, Salinas, CA",
        "Pack Style": "12x1ct Clam",
        "Number of Plants": "240 heads (12 heads x 20 cases)",
        "Lot Code": lot,
        # Preview-only columns must never be written.
        "Confidence": "100%",
        "Status": "Verified",
        "Warnings": "",
        "_Source Path": "C:/x.pdf",
    }
    row.update(overrides)
    return row


def _values(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_new_workbook_gets_headers_and_rows(tmp_path):
    path = tmp_path / "new.xlsx"
    backup = append_rows(str(path), [_row(), _row(so="222222")])
    assert backup == ""
    ws = load_workbook(path).active
    assert ws.title == "Data"
    assert _values(ws) == [
        HEADERS,
        ["August 3, 2026", "123456", "Acme Farms, 1 Main St, Salinas, CA", "12x1ct Clam", "240 heads (12 heads x 20 cases)", "L-77"],
        ["August 3, 2026", "222222", "Acme Farms, 1 Main St, Salinas, CA", "12x1ct Clam", "240 heads (12 heads x 20 cases)", "L-77"],
    ]
    assert ws["A1"].font.bold


def test_reordered_and_extra_columns_stay_aligned(tmp_path):
    path = tmp_path / "custom.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Notes", "Lot Code", "date", "Sales Order Number", "Customer and Location", "Pack Style", "Number  of Plants"])
    ws.append(["keep me", "OLD-1", "July 1, 2026", "111111", "Old Customer", "8x2ct Clam", "16 plants"])
    ws.column_dimensions["A"].width = 33
    wb.save(path)

    append_rows(str(path), [_row()], create_backup_first=False)

    ws = load_workbook(path).active
    assert _values(ws)[2] == [
        None, "L-77", "August 3, 2026", "123456", "Acme Farms, 1 Main St, Salinas, CA", "12x1ct Clam",
        "240 heads (12 heads x 20 cases)",
    ]
    assert ws.max_column == 7
    assert ws.column_dimensions["A"].width == 33


def test_missing_headers_are_added_to_the_right(tmp_path):
    path = tmp_path / "partial.xlsx"
    wb = Workbook()
    wb.active.append(["Date", "Sales Order Number", "Customer and Location", "Comments"])
    wb.save(path)

    append_rows(str(path), [_row()], create_backup_first=False)

    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == ["Date", "Sales Order Number", "Customer and Location", "Comments", "Pack Style", "Number of Plants", "Lot Code"]
    assert [c.value for c in ws[2]] == ["August 3, 2026", "123456", "Acme Farms, 1 Main St, Salinas, CA", None, "12x1ct Clam", "240 heads (12 heads x 20 cases)", "L-77"]


def test_unrecognized_headers_keep_original_fixed_layout(tmp_path):
    path = tmp_path / "renamed.xlsx"
    wb = Workbook()
    wb.active.append(["When", "SO #", "Customer", "Pack", "Plants", "Lot"])
    wb.save(path)

    append_rows(str(path), [_row()], create_backup_first=False)

    ws = load_workbook(path).active
    assert ws.max_column == 6
    assert [c.value for c in ws[2]] == ["August 3, 2026", "123456", "Acme Farms, 1 Main St, Salinas, CA", "12x1ct Clam", "240 heads (12 heads x 20 cases)", "L-77"]


def test_writes_to_data_sheet_even_when_another_sheet_is_active(tmp_path):
    path = tmp_path / "multi.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Monthly summary"
    data = wb.create_sheet("Shipments")
    data.append(HEADERS)
    wb.active = 0
    wb.save(path)

    append_rows(str(path), [_row()], create_backup_first=False)

    wb = load_workbook(path)
    assert _values(wb["Summary"]) == [["Monthly summary"]]
    assert _values(wb["Shipments"])[1][1] == "123456"


def test_styled_blank_rows_do_not_leave_gaps(tmp_path):
    path = tmp_path / "styled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append(["July 1, 2026", "111111", "Old", "8x2ct Clam", "16 plants", "OLD-1"])
    thin = Side(style="thin")
    for row in range(3, 50):
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(top=thin, bottom=thin)
            cell.fill = PatternFill("solid", fgColor="FFFFCC")
    wb.save(path)

    append_rows(str(path), [_row()], create_backup_first=False)

    ws = load_workbook(path).active
    assert ws.cell(row=3, column=2).value == "123456"
    assert ws.cell(row=4, column=2).value is None


def test_backup_created_and_no_temp_file_left(tmp_path):
    path = tmp_path / "book.xlsx"
    append_rows(str(path), [_row()])
    backup = append_rows(str(path), [_row(so="222222")])

    assert backup and os.path.exists(backup)
    assert len(_values(load_workbook(backup).active)) == 2
    assert len(_values(load_workbook(path).active)) == 3
    assert not list(tmp_path.glob("*.saving*"))


def test_locked_workbook_raises_and_is_untouched(tmp_path):
    path = tmp_path / "locked.xlsx"
    append_rows(str(path), [_row()])
    before = path.read_bytes()
    os.chmod(path, stat.S_IREAD)
    try:
        with pytest.raises(PermissionError):
            append_rows(str(path), [_row(so="222222")])
    finally:
        os.chmod(path, stat.S_IREAD | stat.S_IWRITE)
    assert path.read_bytes() == before
    assert not (tmp_path / "Allisons Data Extractor Backups").exists()
    assert not list(tmp_path.glob("*.saving*"))


def test_find_duplicates_in_file_and_batch(tmp_path):
    path = tmp_path / "dupes.xlsx"
    append_rows(str(path), [_row(so="123456", lot="L-77"), _row(so="555555", lot="X")])

    duplicates = find_duplicates(str(path), [
        _row(so="123456", lot="l-77"),
        _row(so="555555", lot="DIFFERENT"),
        _row(so="999999"),
        _row(so="999999"),
    ])

    assert [(d["sales_order"], d["excel_row"], d["match_type"]) for d in duplicates] == [
        ("123456", 2, "Sales Order Number + Lot Code"),
        ("555555", 3, "Sales Order Number"),
        ("999999", "incoming row 3", "Duplicate within current batch"),
    ]


def test_known_values(tmp_path):
    path = tmp_path / "known.xlsx"
    append_rows(str(path), [_row(lot="MUKFI232"), _row(lot="MU01233"), _row(lot="")])
    assert known_values(str(path), "Lot Code") == {"MUKFI232", "MU01233"}
    assert known_values(str(tmp_path / "missing.xlsx"), "Lot Code") == set()


def test_find_duplicates_without_file(tmp_path):
    assert find_duplicates(str(tmp_path / "missing.xlsx"), [_row()]) == []
