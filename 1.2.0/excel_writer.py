"""Excel output helpers for Allison's Data Extractor.

Adds automatic backups and duplicate detection while preserving the original
six-column workbook format.
"""

import os
import shutil
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from pdf_extractor import HEADERS


def _new_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(HEADERS)
    for col_idx in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
    return wb, ws


def _autosize_columns(ws):
    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = len(str(header))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)


def create_backup(filepath: str) -> str:
    """Create a timestamped backup beside an existing workbook."""
    if not os.path.exists(filepath):
        return ""
    folder = os.path.dirname(os.path.abspath(filepath))
    backup_dir = os.path.join(folder, "Allisons Data Extractor Backups")
    os.makedirs(backup_dir, exist_ok=True)
    stem, ext = os.path.splitext(os.path.basename(filepath))
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"{stem}_backup_{stamp}{ext or '.xlsx'}")
    shutil.copy2(filepath, backup_path)
    return backup_path


def _existing_rows(filepath: str):
    if not os.path.exists(filepath):
        return []
    wb = load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        indexes = {name: idx for idx, name in enumerate(headers)}
        rows = []
        for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            item = {h: values[indexes[h]] if h in indexes and indexes[h] < len(values) else "" for h in HEADERS}
            item["_excel_row"] = excel_row
            rows.append(item)
        return rows
    finally:
        wb.close()


def find_duplicates(filepath: str, rows: list) -> list:
    """Find likely duplicate shipments based primarily on Sales Order Number."""
    existing = _existing_rows(filepath)
    duplicates = []
    seen_incoming = {}
    for incoming_index, incoming in enumerate(rows):
        so = str(incoming.get("Sales Order Number", "") or "").strip()
        lot = str(incoming.get("Lot Code", "") or "").strip().upper()
        if not so:
            continue
        if so in seen_incoming:
            duplicates.append({
                "incoming_index": incoming_index,
                "sales_order": so,
                "lot_code": lot,
                "excel_row": f"incoming row {seen_incoming[so] + 1}",
                "match_type": "Duplicate within current batch",
            })
        else:
            seen_incoming[so] = incoming_index
        for old in existing:
            old_so = str(old.get("Sales Order Number", "") or "").strip()
            old_lot = str(old.get("Lot Code", "") or "").strip().upper()
            if old_so != so:
                continue
            match_type = "Sales Order Number"
            if lot and old_lot and lot == old_lot:
                match_type = "Sales Order Number + Lot Code"
            duplicates.append({
                "incoming_index": incoming_index,
                "sales_order": so,
                "lot_code": lot,
                "excel_row": old.get("_excel_row"),
                "match_type": match_type,
            })
            break
    return duplicates


def append_rows(filepath: str, rows: list, create_backup_first: bool = True):
    """Append rows and return the backup path created before writing, if any."""
    backup_path = create_backup(filepath) if create_backup_first else ""

    if os.path.exists(filepath):
        wb = load_workbook(filepath)
        ws = wb.active
        if ws.max_row == 1 and all(c.value is None for c in ws[1]):
            ws.append(HEADERS)
            for col_idx in range(1, len(HEADERS) + 1):
                ws.cell(row=1, column=col_idx).font = Font(bold=True)
    else:
        wb, ws = _new_workbook()

    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])

    _autosize_columns(ws)
    wb.save(filepath)
    return backup_path
